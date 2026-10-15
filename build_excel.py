"""Build the Excel model: Shopee_Monee_Unit_Economics.xlsx.

Every derived number is a live Excel formula that points back to the Data sheet, and every Data row
links to the Sea filing it came from, so a reviewer can trace any figure. Sheets:
README, Data, Shopee_Metrics, Growth_Drivers, Monee, Target_2026, Checks.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
OUT = HERE / "Shopee_Monee_Unit_Economics.xlsx"
df = pd.read_csv(HERE / "data" / "sea_quarterly_clean.csv")
N = len(df)
FIRST, LAST = 2, N + 1  # data rows on every row-aligned sheet

ORANGE, DARK, LIGHT, INPUT = "EE4D2D", "1F1F1F", "FDEDEA", "FFF4CC"
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor=ORANGE)
TITLE = Font(bold=True, size=14, color=ORANGE)
THIN = Border(bottom=Side(style="thin", color="DDDDDD"))

wb = Workbook()


def header(ws, cols, widths=None, row=1):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill = HEAD, FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(c, 14)
    ws.row_dimensions[row].height = 42
    ws.freeze_panes = ws.cell(row=row + 1, column=2)


# ---------------- README ----------------
ws = wb.active
ws.title = "README"
ws.column_dimensions["A"].width = 120
lines = [
    ("Shopee & Monee unit economics, Q1 2019 - Q2 2026", TITLE),
    ("Source: Sea Limited (NYSE: SE) quarterly earnings releases (Form 6-K, exhibit 99.1) on SEC EDGAR. Each Data row links to its filing.", None),
    ("Question: is Shopee's growth becoming more profitable, is Monee's fast loan growth staying healthy, and, before H2 is reported, how likely is the 'US$1 billion Shopee adjusted EBITDA in FY2026' expectation (nowcast as of 15 Oct 2026)?", None),
    ("", None),
    ("How to read this file", Font(bold=True)),
    ("- Yellow cells on Target_2026 are assumptions you can change. Everything else is a formula or a reported number.", None),
    ("- Money in US$ millions unless the column says bn. Take rates and margins are % of GMV unless stated otherwise.", None),
    ("", None),
    ("Definitions", Font(bold=True)),
    ("- GMV: value of goods sold on Shopee. Gross orders: number of orders. Average order value (AOV) = GMV / gross orders.", None),
    ("- Marketplace revenue = core marketplace revenue (transaction fees + advertising) + value-added services (mainly logistics). Shopee's own product sales are excluded.", None),
    ("- Take rate = revenue / GMV. Core take rate uses core marketplace revenue only (reported since Q3 2022).", None),
    ("- Adjusted EBITDA (Sea's segment measure) = operating income + depreciation & amortization. Share-based compensation is unallocated, not in Shopee's figure.", None),
    ("", None),
    ("Data notes and limits", Font(bold=True)),
    ("- Sea prints most bullet figures rounded (e.g. GMV to US$0.1 bn); adjusted EBITDA comes from exact tables (US$ thousands). Ratios inherit that rounding.", None),
    ("- Q1 and Q2 2023: Sea did not report GMV or gross orders. They are derived from the YoY growth Sea reported for Q1/Q2 2024 and flagged; the FY2023 totals confirm them (within 0.7%).", None),
    ("- Q1 and Q2 2020 adjusted EBITDA were later restated by Sea by a few million; this file keeps first-reported values (see Checks).", None),
    ("- Monee's loan measure changed: net on-book receivables (to Q2 2023), gross on-book receivables (Q3 2023), principal outstanding incl. off-book (from Q4 2023). Growth rates are only computed within one basis.", None),
    ("- Before Q3 2023, Sea gave the 90-day NPL ratio only as bounds ('less than 4%', 'around 2%'). The drop to 'less than 2%' in Q4 2022 was mainly from shortening the write-off period from 180 to 120 days; Sea said it would be about 5% without that change.", None),
    ("- From Q3 2026 Sea excludes deferred-revenue effects from adjusted EBITDA. The Q2 2026 reconciliation shows this effect is zero for Shopee (it is Garena only), so the FY2026 test is not affected.", None),
    ("- No country split: Sea does not report Shopee Indonesia separately.", None),
]
for i, (text, font) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=text)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    if font:
        c.font = font

# ---------------- Data ----------------
ws = wb.create_sheet("Data")
cols = [("Period", "period"), ("Release date", "release_date"), ("GMV (US$ bn)", "gmv_usd_bn"), ("GMV source", "gmv_source"),
        ("Gross orders (bn)", "gross_orders_bn"), ("Orders source", "orders_source"), ("GMV YoY reported %", "gmv_yoy_pct"),
        ("Orders YoY reported %", "orders_yoy_pct"), ("Shopee revenue (US$ m)", "shopee_revenue_usd_m"),
        ("Marketplace revenue (US$ m)", "marketplace_usd_m"), ("Core marketplace revenue (US$ m)", "core_marketplace_usd_m"),
        ("Value-added services revenue (US$ m)", "vas_usd_m"), ("Shopee adj. EBITDA (US$ m)", "shopee_adj_ebitda_usd_m"),
        ("Monee revenue (US$ m)", "monee_revenue_usd_m"), ("Monee adj. EBITDA (US$ m)", "monee_adj_ebitda_usd_m"),
        ("Monee loans (US$ bn)", "loans_usd_bn"), ("Loan basis", "loans_basis"), ("NPL 90+ days %", "npl_90d_pct"),
        ("NPL as reported (bounds)", "npl_note"), ("Source filing", "source")]
header(ws, [c for c, _ in cols], {"Loan basis": 34, "Source filing": 18, "GMV source": 20, "Orders source": 20})
C = {key: get_column_letter(i) for i, (_, key) in enumerate(cols, 1)}  # column letter by data key
for r, row in enumerate(df.itertuples(index=False), FIRST):
    rec = row._asdict()
    for i, (_, key) in enumerate(cols, 1):
        v = rec[key]
        v = None if (isinstance(v, float) and pd.isna(v)) else v
        cell = ws.cell(row=r, column=i, value=v)
        if key == "source":
            cell.value, cell.hyperlink, cell.font = "6-K ex.99.1", v, Font(color="0563C1", underline="single")
        if key in ("gmv_source", "orders_source") and v and "derived" in v:
            for k in ("gmv_usd_bn", "gross_orders_bn"):
                ws.cell(row=r, column=list(C).index(k) + 1).fill = PatternFill("solid", fgColor=INPUT)
        cell.border = THIN


def d(key, r):
    return f"Data!{C[key]}{r}"


# ---------------- row-aligned formula sheets ----------------
def formula_sheet(name, spec, widths=None):
    s = wb.create_sheet(name)
    header(s, ["Period"] + [c for c, _, _ in spec], widths)
    for r in range(FIRST, LAST + 1):
        s.cell(row=r, column=1, value=f"=Data!A{r}")
        for j, (_, f, fmt) in enumerate(spec, 2):
            cell = s.cell(row=r, column=j, value=f(r))
            cell.number_format = fmt
            cell.border = THIN
    return s


def safe(expr, *needs):
    cond = ",".join(f'{n}=""' for n in needs)
    return f'=IF(OR({cond}),"",{expr})'


def yoy(key):
    return lambda r: "" if r - 4 < FIRST else safe(f"{d(key, r)}/{d(key, r - 4)}-1", d(key, r), d(key, r - 4))


metrics = [
    ("AOV (US$ per order)", lambda r: safe(f"{d('gmv_usd_bn', r)}/{d('gross_orders_bn', r)}", d('gmv_usd_bn', r), d('gross_orders_bn', r)), "0.00"),
    ("Marketplace take rate", lambda r: safe(f"{d('marketplace_usd_m', r)}/({d('gmv_usd_bn', r)}*1000)", d('marketplace_usd_m', r), d('gmv_usd_bn', r)), "0.0%"),
    ("Core take rate (fees + ads)", lambda r: safe(f"{d('core_marketplace_usd_m', r)}/({d('gmv_usd_bn', r)}*1000)", d('core_marketplace_usd_m', r)), "0.0%"),
    ("VAS (logistics) % GMV", lambda r: safe(f"{d('vas_usd_m', r)}/({d('gmv_usd_bn', r)}*1000)", d('vas_usd_m', r)), "0.0%"),
    ("Adj. EBITDA per order (US$)", lambda r: safe(f"{d('shopee_adj_ebitda_usd_m', r)}/({d('gross_orders_bn', r)}*1000)", d('gross_orders_bn', r)), "0.000"),
    ("Adj. EBITDA % GMV", lambda r: safe(f"{d('shopee_adj_ebitda_usd_m', r)}/({d('gmv_usd_bn', r)}*1000)", d('gmv_usd_bn', r)), "0.00%"),
    ("Adj. EBITDA % marketplace revenue", lambda r: safe(f"{d('shopee_adj_ebitda_usd_m', r)}/{d('marketplace_usd_m', r)}", d('marketplace_usd_m', r)), "0.0%"),
    ("GMV YoY", yoy("gmv_usd_bn"), "0.0%"),
    ("Orders YoY", yoy("gross_orders_bn"), "0.0%"),
    ("Marketplace revenue YoY", yoy("marketplace_usd_m"), "0.0%"),
    ("Trailing 4Q GMV (US$ bn)", lambda r: "" if r - 3 < FIRST else f"=SUM(Data!{C['gmv_usd_bn']}{r - 3}:{C['gmv_usd_bn']}{r})", "0.0"),
    ("Trailing 4Q adj. EBITDA (US$ m)", lambda r: "" if r - 3 < FIRST else f"=SUM(Data!{C['shopee_adj_ebitda_usd_m']}{r - 3}:{C['shopee_adj_ebitda_usd_m']}{r})", "#,##0.0"),
]
sm = formula_sheet("Shopee_Metrics", metrics)


def mref(col_name, r):
    return f"Shopee_Metrics!{get_column_letter(2 + [m[0] for m in metrics].index(col_name))}{r}"


def ln_share(part):
    def f(r):
        if r - 4 < FIRST:
            return ""
        o = f"LN({d('gross_orders_bn', r)}/{d('gross_orders_bn', r - 4)})"
        a = f"LN({mref('AOV (US$ per order)', r)}/{mref('AOV (US$ per order)', r - 4)})"
        t = f"LN({mref('Marketplace take rate', r)}/{mref('Marketplace take rate', r - 4)})"
        expr = {"orders": o, "aov": a, "take": t, "total": f"EXP({o}+{a}+{t})-1"}[part]
        if part != "total":
            expr = f"{expr}/({o}+{a}+{t})"
        return f'=IFERROR({expr},"")'
    return f


drivers = [
    ("Marketplace revenue YoY (= orders x AOV x take rate)", ln_share("total"), "0.0%"),
    ("Share from more orders", ln_share("orders"), "0%"),
    ("Share from AOV", ln_share("aov"), "0%"),
    ("Share from higher take rate", ln_share("take"), "0%"),
]
gd = formula_sheet("Growth_Drivers", drivers, {"Marketplace revenue YoY (= orders x AOV x take rate)": 22})
gd.cell(row=LAST + 2, column=1, value="Method: log decomposition of YoY marketplace revenue growth, ln(R/R-4) = ln(orders ratio) + ln(AOV ratio) + ln(take-rate ratio); each share = its log term / total log growth.").font = Font(italic=True, color="666666")

monee = [
    ("Monee revenue (US$ m)", lambda r: f"=IF({d('monee_revenue_usd_m', r)}=\"\",\"\",{d('monee_revenue_usd_m', r)})", "#,##0.0"),
    ("Monee adj. EBITDA (US$ m)", lambda r: f"={d('monee_adj_ebitda_usd_m', r)}", "#,##0.0"),
    ("Adj. EBITDA margin", lambda r: safe(f"{d('monee_adj_ebitda_usd_m', r)}/{d('monee_revenue_usd_m', r)}", d('monee_revenue_usd_m', r)), "0.0%"),
    ("Loans (US$ bn)", lambda r: f"=IF({d('loans_usd_bn', r)}=\"\",\"\",{d('loans_usd_bn', r)})", "0.0"),
    ("Loan basis", lambda r: f"=IF({d('loans_basis', r)}=\"\",\"\",{d('loans_basis', r)})", "@"),
    ("Loan YoY (same basis only)", lambda r: "" if r - 4 < FIRST else
        f'=IF(OR({d("loans_usd_bn", r)}="",{d("loans_usd_bn", r - 4)}="",{d("loans_basis", r)}<>{d("loans_basis", r - 4)}),"",{d("loans_usd_bn", r)}/{d("loans_usd_bn", r - 4)}-1)', "0.0%"),
    ("NPL 90+ days %", lambda r: f"=IF({d('npl_90d_pct', r)}=\"\",\"\",{d('npl_90d_pct', r)}/100)", "0.0%"),
    ("NPL as reported (bounds)", lambda r: f"=IF({d('npl_note', r)}=\"\",\"\",{d('npl_note', r)})", "@"),
    ("Approx. NPL amount (US$ m)", lambda r: safe(f"{d('npl_90d_pct', r)}/100*{d('loans_usd_bn', r)}*1000", d('npl_90d_pct', r), d('loans_usd_bn', r)), "#,##0"),
    ("NPL vs loans 4Q earlier (seasoning view)", lambda r: "" if r - 4 < FIRST else
        f'=IF(OR({d("npl_90d_pct", r)}="",{d("loans_usd_bn", r - 4)}="",{d("loans_basis", r)}<>{d("loans_basis", r - 4)}),"",{d("npl_90d_pct", r)}/100*{d("loans_usd_bn", r)}/{d("loans_usd_bn", r - 4)})', "0.0%"),
]
mo = formula_sheet("Monee", monee, {"Loan basis": 34, "NPL vs loans 4Q earlier (seasoning view)": 20})
mo.cell(row=LAST + 2, column=1, value="Seasoning view: when a loan book grows fast, new loans have not had time to go bad, which flatters the NPL ratio. Dividing today's NPL amount by the book four quarters earlier shows the ratio without that effect. Approximate: the NPL % is printed to one decimal.").font = Font(italic=True, color="666666")

# ---------------- Target_2026 ----------------
ws = wb.create_sheet("Target_2026")
ws.column_dimensions["A"].width = 46
for col in "BCDE":
    ws.column_dimensions[col].width = 16
row_of = {p: r for r, p in enumerate(df["period"], FIRST)}
q = lambda key, p: d(key, row_of[p])
ws["A1"], ws["A1"].font = "Nowcast: can Shopee reach US$1 billion adjusted EBITDA in FY2026?", TITLE
ws["A2"] = "Sea's CEO said in the Q2 2026 release that he is 'optimistic' Shopee will reach US$1 billion adjusted EBITDA for the full year."
ws["A3"] = "Pre-results estimate (nowcast) made on 15 Oct 2026 from H1 2026 data, before Sea reports Q3 2026 (mid-November 2026) and FY2026 (around March 2027)."
ws["A3"].font = Font(italic=True, color="666666")
ws["A4"], ws["A4"].font = "Actuals (from Data)", Font(bold=True)
facts = [("Q1 2026 adj. EBITDA (US$ m)", f"={q('shopee_adj_ebitda_usd_m', '2026Q1')}", "#,##0.0"),
         ("Q2 2026 adj. EBITDA (US$ m)", f"={q('shopee_adj_ebitda_usd_m', '2026Q2')}", "#,##0.0"),
         ("H1 2026 adj. EBITDA (US$ m)", "=B5+B6", "#,##0.0"),
         ("Needed in H2 2026 for US$1 bn (US$ m)", "=1000-B7", "#,##0.0"),
         ("Q3 2025 GMV (US$ bn)", f"={q('gmv_usd_bn', '2025Q3')}", "0.0"),
         ("Q4 2025 GMV (US$ bn)", f"={q('gmv_usd_bn', '2025Q4')}", "0.0"),
         ("H1 2026 GMV growth YoY", f"=({q('gmv_usd_bn', '2026Q1')}+{q('gmv_usd_bn', '2026Q2')})/({q('gmv_usd_bn', '2025Q1')}+{q('gmv_usd_bn', '2025Q2')})-1", "0.0%"),
         ("H2 2025 adj. EBITDA % GMV", f"=({q('shopee_adj_ebitda_usd_m', '2025Q3')}+{q('shopee_adj_ebitda_usd_m', '2025Q4')})/(({q('gmv_usd_bn', '2025Q3')}+{q('gmv_usd_bn', '2025Q4')})*1000)", "0.00%"),
         ("H1 2026 adj. EBITDA % GMV", f"=B7/(({q('gmv_usd_bn', '2026Q1')}+{q('gmv_usd_bn', '2026Q2')})*1000)", "0.00%"),
         ("Q2 2026 adj. EBITDA % GMV", f"={q('shopee_adj_ebitda_usd_m', '2026Q2')}/({q('gmv_usd_bn', '2026Q2')}*1000)", "0.00%"),
         ("H2 2025 EBITDA as % of H1 2025 EBITDA (seasonality)", f"=({q('shopee_adj_ebitda_usd_m', '2025Q3')}+{q('shopee_adj_ebitda_usd_m', '2025Q4')})/({q('shopee_adj_ebitda_usd_m', '2025Q1')}+{q('shopee_adj_ebitda_usd_m', '2025Q2')})", "0%")]
for i, (label, f, fmt) in enumerate(facts, 5):
    ws.cell(row=i, column=1, value=label)
    c = ws.cell(row=i, column=2, value=f)
    c.number_format = fmt
# rows 5..15 used
ws["A17"], ws["A17"].font = "Scenarios for H2 2026 (yellow = assumption)", Font(bold=True)
for j, name in enumerate(["Bear", "Base", "Bull"], 2):
    c = ws.cell(row=18, column=j, value=name)
    c.font, c.fill = HEAD, FILL
assumptions = [
    ("H2 GMV growth YoY", ["=B11-0.05", "=B11", "=B11+0.03"], "0.0%"),
    ("H2 adj. EBITDA % GMV", ["=B12", "=B13", "=B14"], "0.00%"),
]
for i, (label, vals, fmt) in enumerate(assumptions, 19):
    ws.cell(row=i, column=1, value=label)
    for j, v in enumerate(vals, 2):
        c = ws.cell(row=i, column=j, value=v)
        c.number_format, c.fill = fmt, PatternFill("solid", fgColor=INPUT)
ws["A21"] = "Why these: Bear = growth 5 pts slower and H2 margin like H2 2025 (mega-sale season); Base = H1 2026 growth and margin; Bull = 3 pts faster growth at the Q2 2026 margin."
ws["A21"].font = Font(italic=True, color="666666")
outputs = [
    ("H2 2026 GMV (US$ bn)", "=($B$9+$B$10)*(1+{c}19)", "0.0"),
    ("H2 2026 adj. EBITDA (US$ m)", "={c}23*1000*{c}20", "#,##0.0"),
    ("FY2026 adj. EBITDA (US$ m)", "=$B$7+{c}24", "#,##0.0"),
    ("Gap to US$1 bn (US$ m)", "={c}25-1000", "#,##0.0"),
    ("Reaches US$1 bn?", '=IF({c}25>=1000,"Yes","No")', "@"),
    ("H2 margin needed for US$1 bn (% GMV)", "=$B$8/({c}23*1000)", "0.00%"),
]
for i, (label, f, fmt) in enumerate(outputs, 23):
    ws.cell(row=i, column=1, value=label).font = Font(bold=i in (25, 27))
    for j, col in enumerate("BCD", 2):
        c = ws.cell(row=i, column=j, value=f.format(c=col))
        c.number_format = fmt
        c.alignment = Alignment(horizontal="right")
ws["A30"], ws["A30"].font = "Simple seasonal check", Font(bold=True)
ws["A31"] = "FY2026 if H2 repeats 2025's H2/H1 ratio (US$ m)"
ws["B31"], ws["B31"].number_format = "=B7*(1+B15)", "#,##0.0"
ws["A33"] = "Check after Sea's Q3 2026 release (mid-November): put actual Q3 adj. EBITDA and GMV here and compare with the Base case."
ws["A33"].font = Font(italic=True, color="666666")
ws["A34"], ws["A35"] = "Actual Q3 2026 adj. EBITDA (US$ m)", "Actual Q3 2026 GMV (US$ bn)"
for cell in ("B34", "B35"):
    ws[cell].fill = PatternFill("solid", fgColor=INPUT)

# ---------------- Checks ----------------
ws = wb.create_sheet("Checks")
ws.column_dimensions["A"].width = 130
for i, line in enumerate((HERE / "data" / "checks.md").read_text(encoding="utf-8").splitlines(), 1):
    cell = ws.cell(row=i, column=1, value=line.lstrip("# "))
    if line.startswith("#"):  # markdown heading -> bold text (a leading '#' looks like an Excel error)
        cell.font = TITLE if line.startswith("# ") else Font(bold=True)

# ---------------- charts ----------------
def line_chart(sheet, title, cols, anchor, y_fmt="0%", start=FIRST):
    """start: first data row to plot; blank formula cells ("") would otherwise plot as zero."""
    ch = LineChart()
    ch.title, ch.height, ch.width = title, 8, 18
    for col in cols:
        ch.add_data(Reference(sheet, min_col=col, min_row=start, max_row=LAST), titles_from_data=False)
    for s, col in zip(ch.series, cols):
        s.tx = SeriesLabel(v=sheet.cell(row=1, column=col).value)
        s.smooth = False
    ch.set_categories(Reference(sheet, min_col=1, min_row=start, max_row=LAST))
    ch.y_axis.number_format = y_fmt
    ch.x_axis.delete = ch.y_axis.delete = False  # newer Excel hides axes unless this is explicit
    ch.varyColors = False  # one colour per series, not per point
    for s, colour in zip(ch.series, (ORANGE, "555555")):
        s.graphicalProperties.line.solidFill = colour
    ch.legend = ch.legend if len(cols) > 1 else None
    if ch.legend:
        ch.legend.position = "t"
    sheet.add_chart(ch, anchor)


names = [m[0] for m in metrics]
line_chart(sm, "Shopee take rate: fees + ads vs logistics (% GMV)", [2 + names.index("Core take rate (fees + ads)"), 2 + names.index("VAS (logistics) % GMV")], "P2",
           start=row_of["2022Q3"])  # core/VAS split reported from Q3 2022
line_chart(sm, "Shopee adjusted EBITDA % GMV", [2 + names.index("Adj. EBITDA % GMV")], "P20", "0.0%")
bc = BarChart()
bc.title, bc.height, bc.width = "Monee loans (US$ bn) - measure changes in 2023Q3 and 2023Q4, see Loan basis", 8, 18
bc.add_data(Reference(mo, min_col=5, min_row=row_of["2022Q3"], max_row=LAST), titles_from_data=False)
bc.set_categories(Reference(mo, min_col=1, min_row=row_of["2022Q3"], max_row=LAST))
bc.varyColors = False
bc.series[0].graphicalProperties.solidFill = ORANGE
bc.x_axis.delete = bc.y_axis.delete = False
bc.legend = None
mo.add_chart(bc, "M2")

wb.save(OUT)
print("saved", OUT)
